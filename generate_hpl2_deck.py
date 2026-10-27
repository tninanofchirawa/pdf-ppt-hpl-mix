"""
HPL 2.0 - Official Mega Player Auction Presentation Generator
============================================================
A robust, standalone script that reads cricket player records and generates
a high-impact widescreen (16:9) PowerPoint presentation deck for the HPL 2.0 auction.

Author: Antigravity (Specialist in python-pptx presentation generation)
"""

import os
import sys
import re
import argparse
from typing import List, Dict, Any, Union, Optional
import pandas as pd

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ==============================================================================
# DESIGN & COLOR PALETTE CONSTANTS
# ==============================================================================
SLIDE_WIDTH_INCHES = 13.333
SLIDE_HEIGHT_INCHES = 7.5

# Theme Colors
COLOR_BG = RGBColor(11, 15, 25)               # Deep Navy Dark (#0B0F19)
COLOR_CARD_FILL = RGBColor(22, 31, 48)         # Dark Slate Blue (#161F30)
COLOR_CARD_BORDER = RGBColor(39, 53, 79)       # Subtle Border (#27354F)
COLOR_GOLD_ACCENT = RGBColor(250, 204, 21)     # Gold Accent (#FACC15)
COLOR_PRICE_GREEN = RGBColor(34, 197, 94)      # Electric Green (#22C55E)
COLOR_SECONDARY_TEXT = RGBColor(156, 163, 175) # Light Gray (#9CA3AF)
COLOR_WHITE_TEXT = RGBColor(248, 249, 250)     # White (#F8F9FA)

# Badge Colors
COLOR_ROLE_BG = RGBColor(30, 41, 59)          # Slate surface for role
COLOR_ROLE_BORDER = RGBColor(59, 130, 246)     # Electric Blue accent
COLOR_COUNTRY_INDIA_BG = RGBColor(16, 58, 42)  # Dark Forest Green
COLOR_COUNTRY_INDIA_BORDER = RGBColor(16, 185, 129) # Emerald Green
COLOR_COUNTRY_OVS_BG = RGBColor(51, 35, 78)    # Dark Purple
COLOR_COUNTRY_OVS_BORDER = RGBColor(168, 85, 247) # Violet

# Generalized Cover Slide Defaults
DEFAULT_LEAGUE_TAG = "HPL PREMIER LEAGUE"
DEFAULT_AUCTION_TITLE = "HPL 2.0"
DEFAULT_AUCTION_SUBTITLE = "OFFICIAL MEGA PLAYER AUCTION 2026"
DEFAULT_EVENT_TAG = "LIVE SQUAD SELECTION & BIDDING"
DEFAULT_ORGANIZER = "Organized by Tournament Committee"


# ==============================================================================
# SAMPLE MARQUEE DATASET (FOR INSTANT PREVIEWS & FALLBACKS)
# ==============================================================================
SAMPLE_PLAYERS: List[Dict[str, Any]] = [
    {
        "Player_Name": "VIRAT KOHLI",
        "Role": "BATSMAN",
        "Country": "INDIA",
        "Points": 9.8,
        "Base_Price_Cr": 2.00,
        "Set_Name": "SET 1: MARQUEE BATSMEN",
        "Image_Path": "outputs/images_current/virat_kohli.png"
    },
    {
        "Player_Name": "JASPRIT BUMRAH",
        "Role": "BOWLER",
        "Country": "INDIA",
        "Points": 9.9,
        "Base_Price_Cr": 2.00,
        "Set_Name": "SET 2: MARQUEE BOWLERS",
        "Image_Path": "outputs/images_current/jasprit_bumrah.png"
    },
    {
        "Player_Name": "TRAVIS HEAD",
        "Role": "BATSMAN",
        "Country": "OVERSEAS",
        "Points": 9.2,
        "Base_Price_Cr": 2.00,
        "Set_Name": "SET 1: MARQUEE BATSMEN",
        "Image_Path": "outputs/images_current/travis_head.png"
    },
    {
        "Player_Name": "KL RAHUL",
        "Role": "WICKETKEEPER",
        "Country": "INDIA",
        "Points": 9.0,
        "Base_Price_Cr": 2.00,
        "Set_Name": "SET 3: WICKETKEEPERS",
        "Image_Path": "outputs/images_current/kl_rahul_wk.png"
    },
    {
        "Player_Name": "HARDIK PANDYA",
        "Role": "ALL-ROUNDER",
        "Country": "INDIA",
        "Points": 9.4,
        "Base_Price_Cr": 2.00,
        "Set_Name": "SET 4: ALL-ROUNDERS",
        "Image_Path": "outputs/images_current/hardik_pandya.png"
    },
    {
        "Player_Name": "MS DHONI",
        "Role": "WICKETKEEPER",
        "Country": "INDIA",
        "Points": 10.0,
        "Base_Price_Cr": 2.00,
        "Set_Name": "SET 3: WICKETKEEPERS",
        "Image_Path": "outputs/images_current/ms_dhoni_wk.png"
    }
]


# ==============================================================================
# DATA NORMALIZATION & SANITIZATION HELPERS
# ==============================================================================
def normalize_role(role_raw: Any) -> str:
    """Normalizes role strings to standard auction badges."""
    s = str(role_raw).strip().upper()
    if any(k in s for k in ["WICKET", "WK"]):
        return "WICKETKEEPER"
    elif any(k in s for k in ["BOWL", "PACER", "SPINNER"]):
        return "BOWLER"
    elif any(k in s for k in ["ROUNDER", "ALL-ROUND"]):
        return "ALL-ROUNDER"
    elif any(k in s for k in ["BAT", "BATSMAN", "BATTER"]):
        return "BATSMAN"
    return "PLAYER"


def normalize_country(country_raw: Any) -> str:
    """Normalizes nationality tag to INDIA or OVERSEAS."""
    s = str(country_raw).strip().upper()
    if s in ["INDIA", "DOMESTIC", "IND"]:
        return "INDIA"
    return "OVERSEAS"


def normalize_points(points_raw: Any) -> str:
    """Formats points/rating to X.X string."""
    try:
        val = float(points_raw)
        return f"{val:.1f}"
    except (ValueError, TypeError):
        s = str(points_raw).strip()
        return s if s and s != "nan" else "8.0"


def normalize_base_price(price_raw: Any, points_raw: Any) -> str:
    """Formats base price as ₹X.XX Cr with smart fallback based on rating points."""
    s = str(price_raw).strip()
    if "CR" in s.upper():
        clean = re.sub(r'[^0-9\.]', '', s)
        try:
            val = float(clean)
            return f"₹ {val:.2f} Cr"
        except ValueError:
            return s

    try:
        val = float(s)
        return f"₹ {val:.2f} Cr"
    except (ValueError, TypeError):
        pass

    # Dynamic fallback based on player rating points
    try:
        pts = float(points_raw)
    except (ValueError, TypeError):
        pts = 7.0

    if pts >= 9.0:
        return "₹ 2.00 Cr"
    elif pts >= 8.0:
        return "₹ 1.50 Cr"
    elif pts >= 7.0:
        return "₹ 1.00 Cr"
    elif pts >= 6.0:
        return "₹ 0.50 Cr"
    else:
        return "₹ 0.30 Cr"


def load_dataset(source: Union[str, List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    """
    Loads player data from a list of dicts, CSV file, or Excel file.
    Maps heterogeneous column names to standard keys.
    """
    if isinstance(source, list):
        raw_records = source
    elif isinstance(source, str):
        if source.strip().upper() == "SAMPLE":
            return SAMPLE_PLAYERS

        if not os.path.exists(source):
            alt_candidates = [
                os.path.join("outputs", os.path.basename(source)),
                os.path.join("inputs", os.path.basename(source)),
                os.path.basename(source)
            ]
            found = False
            for alt in alt_candidates:
                if os.path.exists(alt):
                    source = alt
                    found = True
                    break
            if not found:
                print(f"[!] Source dataset '{source}' not found. Using sample marquee players.")
                return SAMPLE_PLAYERS

        if source.lower().endswith(".csv"):
            df = pd.read_csv(source)
        else:
            df = pd.read_excel(source)
        raw_records = df.to_dict(orient="records")
    else:
        return SAMPLE_PLAYERS

    normalized_players: List[Dict[str, Any]] = []

    for idx, rec in enumerate(raw_records):
        p_name = rec.get("Player_Name") or rec.get("Name") or rec.get("player_name") or f"Player {idx + 1}"
        role = rec.get("Role") or rec.get("Category") or rec.get("role") or rec.get("category") or "BATSMAN"
        country = rec.get("Country") or rec.get("country") or rec.get("Nationality") or "INDIA"
        points = rec.get("Points") or rec.get("Rating_Points") or rec.get("Rating") or rec.get("rating") or "8.0"
        price = rec.get("Base_Price_Cr") or rec.get("Base_Price") or rec.get("base_price_cr") or "-"
        img = rec.get("Image_Path") or rec.get("image_path") or rec.get("Photo") or ""
        set_name = rec.get("Set_Name") or rec.get("Set") or ""

        norm_name = str(p_name).strip().upper()
        norm_role = normalize_role(role)
        norm_country = normalize_country(country)
        norm_points = normalize_points(points)
        norm_price = normalize_base_price(price, norm_points)

        if not set_name:
            set_name = f"SET: {norm_role}"

        # Resolve local photo path
        resolved_img = str(img).strip()
        if resolved_img and not os.path.exists(resolved_img):
            cand_paths = [
                os.path.join("outputs", resolved_img),
                os.path.join("outputs", "images_current", os.path.basename(resolved_img)),
                os.path.join("images_current", os.path.basename(resolved_img))
            ]
            for c in cand_paths:
                if os.path.exists(c):
                    resolved_img = c
                    break

        normalized_players.append({
            "Player_Name": norm_name,
            "Role": norm_role,
            "Country": norm_country,
            "Points": norm_points,
            "Base_Price_Cr": norm_price,
            "Set_Name": str(set_name).upper(),
            "Image_Path": resolved_img
        })

    return normalized_players


# ==============================================================================
# SLIDE BUILDERS (COVER & PLAYER AUCTION CARDS)
# ==============================================================================
def set_slide_background(slide, prs: Presentation) -> None:
    """Fills slide background canvas with Deep Navy Dark (#0B0F19)."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()


def build_cover_slide(prs: Presentation,
                      tagline: str = DEFAULT_LEAGUE_TAG,
                      title: str = DEFAULT_AUCTION_TITLE,
                      subtitle: str = DEFAULT_AUCTION_SUBTITLE,
                      event_tag: str = DEFAULT_EVENT_TAG,
                      organizer: str = DEFAULT_ORGANIZER) -> None:
    """
    Builds Slide 1: Front / Cover Slide for 'HPL 2.0'.
    - Large bold title: e.g. 'HPL 2.0' (Gold, 60pt)
    - Subtitle: e.g. 'OFFICIAL MEGA PLAYER AUCTION 2026' (White, 28pt)
    - Tagline: e.g. 'HPL PREMIER LEAGUE'
    - Event Badge: e.g. 'LIVE SQUAD SELECTION & BIDDING'
    - Footer text: e.g. 'Organized by Tournament Committee' (Light Gray, 18pt)
    """
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    set_slide_background(slide, prs)

    # Main Card Container
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.2), Inches(0.8),
        Inches(10.933), Inches(5.9)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_FILL
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(2.0)

    # Top Tagline Pill
    tag_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4.666), Inches(1.3),
        Inches(4.0), Inches(0.55)
    )
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
    tag_box.line.color.rgb = COLOR_GOLD_ACCENT
    tag_box.line.width = Pt(1.5)

    tf_tag = tag_box.text_frame
    tf_tag.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tagline
    p_tag.alignment = PP_ALIGN.CENTER
    p_tag.font.size = Pt(13)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_GOLD_ACCENT

    # Main Title: HPL 2.0
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.1), Inches(10.333), Inches(1.3))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.alignment = PP_ALIGN.CENTER
    p_t.font.size = Pt(60)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_GOLD_ACCENT

    # Subtitle: OFFICIAL MEGA PLAYER AUCTION 2026
    sub_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.4), Inches(10.333), Inches(0.8))
    tf_s = sub_box.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.text = subtitle
    p_s.alignment = PP_ALIGN.CENTER
    p_s.font.size = Pt(28)
    p_s.font.bold = True
    p_s.font.color.rgb = COLOR_WHITE_TEXT

    # Decorative Accent Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4.666), Inches(4.4),
        Inches(4.0), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_PRICE_GREEN
    line.line.fill.background()

    # Sub-badge: LIVE BIDDING EVENT
    event_pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4.416), Inches(4.7),
        Inches(4.5), Inches(0.55)
    )
    event_pill.fill.solid()
    event_pill.fill.fore_color.rgb = RGBColor(16, 58, 42)
    event_pill.line.color.rgb = COLOR_PRICE_GREEN
    event_pill.line.width = Pt(1.5)

    tf_e = event_pill.text_frame
    tf_e.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_e = tf_e.paragraphs[0]
    p_e.text = event_tag
    p_e.alignment = PP_ALIGN.CENTER
    p_e.font.size = Pt(14)
    p_e.font.bold = True
    p_e.font.color.rgb = COLOR_PRICE_GREEN

    # Footer Text
    footer_box = slide.shapes.add_textbox(Inches(1.5), Inches(5.6), Inches(10.333), Inches(0.8))
    tf_f = footer_box.text_frame
    tf_f.word_wrap = True
    p_f = tf_f.paragraphs[0]
    p_f.text = organizer
    p_f.alignment = PP_ALIGN.CENTER
    p_f.font.size = Pt(18)
    p_f.font.color.rgb = COLOR_SECONDARY_TEXT


def build_player_card_slide(prs: Presentation, player: Dict[str, Any]) -> None:
    """
    Builds Slide 2+: Individual Player Auction Card.
    - Zero detailed statistical tables or stat boxes.
    - Header Bar: Category / Set Name (Gold, 18pt bold).
    - Photo Frame: Rounded rectangle (0.8" left, 1.5" top, 4.0" width, 4.8" height).
    - Main Player Info Card:
      - Player Name: Large bold white text (40pt).
      - Role Badge: Colored pill container.
      - Country Tag: Pill container ("INDIA" or "OVERSEAS").
      - Player Rating Badge: Large text "RATING: X.X / 10" in Gold accent.
    - Designated Base Price Box:
      - Prominent container at bottom-right.
      - Text: "BASE PRICE: ₹X.XX Cr" (Electric Green, 32pt bold).
    """
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    set_slide_background(slide, prs)

    # --------------------------------------------------------------------------
    # 1. Header Bar: Category / Set Name
    # --------------------------------------------------------------------------
    set_name = player.get("Set_Name", "SET 1: MARQUEE PLAYERS").upper()
    header_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(0.45),
        Inches(11.733), Inches(0.65)
    )
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = COLOR_CARD_FILL
    header_shape.line.color.rgb = COLOR_CARD_BORDER
    header_shape.line.width = Pt(1.5)

    tf_h = header_shape.text_frame
    tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_h = tf_h.paragraphs[0]
    p_h.text = set_name
    p_h.alignment = PP_ALIGN.CENTER
    p_h.font.size = Pt(18)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_GOLD_ACCENT

    # --------------------------------------------------------------------------
    # 2. Photo Frame (Left Side): 0.8" left, 1.5" top, 4.0" width, 4.8" height
    # --------------------------------------------------------------------------
    photo_frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(1.5),
        Inches(4.0), Inches(4.8)
    )
    photo_frame.fill.solid()
    photo_frame.fill.fore_color.rgb = COLOR_CARD_FILL
    photo_frame.line.color.rgb = COLOR_CARD_BORDER
    photo_frame.line.width = Pt(1.5)

    img_path = str(player.get("Image_Path", "")).strip()
    image_attached = False

    if img_path and os.path.exists(img_path):
        try:
            slide.shapes.add_picture(
                img_path,
                Inches(0.9), Inches(1.6),
                width=Inches(3.8), height=Inches(4.6)
            )
            image_attached = True
        except Exception:
            image_attached = False

    if not image_attached:
        tf_pf = photo_frame.text_frame
        tf_pf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_pf = tf_pf.paragraphs[0]
        p_pf.text = "PHOTO GOES HERE"
        p_pf.alignment = PP_ALIGN.CENTER
        p_pf.font.size = Pt(16)
        p_pf.font.color.rgb = COLOR_SECONDARY_TEXT

    # --------------------------------------------------------------------------
    # 3. Main Player Info Card (Right Side)
    # --------------------------------------------------------------------------
    info_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.1), Inches(1.5),
        Inches(7.433), Inches(3.35)
    )
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = COLOR_CARD_FILL
    info_card.line.color.rgb = COLOR_CARD_BORDER
    info_card.line.width = Pt(1.5)

    # 3a. Player Name: 40pt bold white (with dynamic auto-scaling)
    player_name = str(player.get("Player_Name", "UNKNOWN PLAYER")).upper()
    font_size_pt = 40 if len(player_name) <= 16 else (32 if len(player_name) <= 22 else 26)

    name_box = slide.shapes.add_textbox(Inches(5.35), Inches(1.65), Inches(6.9), Inches(1.15))
    tf_n = name_box.text_frame
    tf_n.word_wrap = True
    p_n = tf_n.paragraphs[0]
    p_n.text = player_name
    p_n.font.size = Pt(font_size_pt)
    p_n.font.bold = True
    p_n.font.color.rgb = COLOR_WHITE_TEXT

    # 3b. Role Badge Container
    role_text = str(player.get("Role", "BATSMAN")).upper()
    role_badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.35), Inches(2.9),
        Inches(2.5), Inches(0.55)
    )
    role_badge.fill.solid()
    role_badge.fill.fore_color.rgb = COLOR_ROLE_BG
    role_badge.line.color.rgb = COLOR_ROLE_BORDER
    role_badge.line.width = Pt(1.5)

    tf_r = role_badge.text_frame
    tf_r.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_r = tf_r.paragraphs[0]
    p_r.text = role_text
    p_r.alignment = PP_ALIGN.CENTER
    p_r.font.size = Pt(14)
    p_r.font.bold = True
    p_r.font.color.rgb = COLOR_WHITE_TEXT

    # 3c. Country/Nationality Tag
    country_text = str(player.get("Country", "INDIA")).upper()
    is_india = (country_text == "INDIA")

    country_badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(8.05), Inches(2.9),
        Inches(2.1), Inches(0.55)
    )
    country_badge.fill.solid()
    country_badge.fill.fore_color.rgb = COLOR_COUNTRY_INDIA_BG if is_india else COLOR_COUNTRY_OVS_BG
    country_badge.line.color.rgb = COLOR_COUNTRY_INDIA_BORDER if is_india else COLOR_COUNTRY_OVS_BORDER
    country_badge.line.width = Pt(1.5)

    tf_c = country_badge.text_frame
    tf_c.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_c = tf_c.paragraphs[0]
    p_c.text = country_text
    p_c.alignment = PP_ALIGN.CENTER
    p_c.font.size = Pt(14)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_WHITE_TEXT

    # 3d. Player Rating Badge (Gold Accent)
    points_text = str(player.get("Points", "8.0")).strip()
    rating_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.35), Inches(3.65),
        Inches(4.8), Inches(0.9)
    )
    rating_box.fill.solid()
    rating_box.fill.fore_color.rgb = COLOR_CARD_FILL
    rating_box.line.color.rgb = COLOR_GOLD_ACCENT
    rating_box.line.width = Pt(1.5)

    tf_rt = rating_box.text_frame
    tf_rt.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_rt = tf_rt.paragraphs[0]
    p_rt.text = f"RATING: {points_text} / 10"
    p_rt.alignment = PP_ALIGN.CENTER
    p_rt.font.size = Pt(22)
    p_rt.font.bold = True
    p_rt.font.color.rgb = COLOR_GOLD_ACCENT

    # --------------------------------------------------------------------------
    # 4. Designated Base Price Box (Bottom Right)
    # --------------------------------------------------------------------------
    base_price_text = str(player.get("Base_Price_Cr", "₹ 2.00 Cr")).strip()
    price_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.1), Inches(5.05),
        Inches(7.433), Inches(1.55)
    )
    price_box.fill.solid()
    price_box.fill.fore_color.rgb = COLOR_CARD_FILL
    price_box.line.color.rgb = COLOR_PRICE_GREEN
    price_box.line.width = Pt(3.0)

    tf_pr = price_box.text_frame
    tf_pr.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_pr = tf_pr.paragraphs[0]
    p_pr.text = f"BASE PRICE: {base_price_text}"
    p_pr.alignment = PP_ALIGN.CENTER
    p_pr.font.size = Pt(32)
    p_pr.font.bold = True
    p_pr.font.color.rgb = COLOR_PRICE_GREEN


# ==============================================================================
# MAIN GENERATOR FUNCTION
# ==============================================================================
def generate_hpl2_presentation(players_data: Union[str, List[Dict[str, Any]]],
                               output_path: str = "HPL_2_Auction_Deck.pptx",
                               tagline: str = DEFAULT_LEAGUE_TAG,
                               title: str = DEFAULT_AUCTION_TITLE,
                               subtitle: str = DEFAULT_AUCTION_SUBTITLE,
                               event_tag: str = DEFAULT_EVENT_TAG,
                               organizer: str = DEFAULT_ORGANIZER) -> str:
    """
    Generates the complete HPL 2.0 Mega Player Auction presentation deck.
    Includes Slide 1 (Cover Slide) and Slides 2+ (Player Auction Cards).
    """
    players = load_dataset(players_data)
    print(f"[*] Initializing HPL 2.0 Auction Presentation ({len(players)} players)...")

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_INCHES)
    prs.slide_height = Inches(SLIDE_HEIGHT_INCHES)

    # 1. Slide 1: Cover Slide
    print("  -> Creating Front / Cover Slide...")
    build_cover_slide(
        prs,
        tagline=tagline,
        title=title,
        subtitle=subtitle,
        event_tag=event_tag,
        organizer=organizer
    )

    # 2. Slides 2+: Player Auction Cards
    print(f"  -> Building {len(players)} Player Auction Cards (No Stats Tables)...")
    for idx, player in enumerate(players):
        build_player_card_slide(prs, player)
        if (idx + 1) % 50 == 0 or (idx + 1) == len(players):
            print(f"     Rendered [{idx + 1}/{len(players)}] cards...")

    # Ensure parent directory exists if path includes subdirectories
    output_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(output_dir, exist_ok=True)

    prs.save(output_path)
    file_size_mb = os.path.getsize(output_path) / (1024 * 1024)
    print(f"[+] HPL 2.0 Auction Presentation saved successfully!")
    print(f"    File: '{output_path}'")
    print(f"    Slides: {len(prs.slides)}")
    print(f"    Size: {file_size_mb:.2f} MB")

    # Also make a mirror copy in outputs/ if generated at root
    if os.path.dirname(output_path) == "":
        outputs_copy = os.path.join("outputs", os.path.basename(output_path))
        try:
            os.makedirs("outputs", exist_ok=True)
            prs.save(outputs_copy)
        except Exception:
            pass

    return output_path


# ==============================================================================
# COMMAND LINE INTERFACE (CLI)
# ==============================================================================
if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="HPL 2.0 - Official Mega Player Auction Presentation Generator",
        formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument(
        "-i", "--input",
        default="outputs/Master_Cricket_Auction.xlsx",
        help="Path to CSV or Excel file containing player records.\n(Default: outputs/Master_Cricket_Auction.xlsx)"
    )
    parser.add_argument(
        "-o", "--output",
        default="HPL_2_Auction_Deck.pptx",
        help="Output presentation path.\n(Default: HPL_2_Auction_Deck.pptx)"
    )
    parser.add_argument(
        "--sample",
        action="store_true",
        help="Generate presentation using built-in marquee sample players."
    )
    parser.add_argument(
        "--title",
        default=DEFAULT_AUCTION_TITLE,
        help=f"Main league title on cover slide (Default: '{DEFAULT_AUCTION_TITLE}')"
    )
    parser.add_argument(
        "--subtitle",
        default=DEFAULT_AUCTION_SUBTITLE,
        help=f"Subtitle on cover slide (Default: '{DEFAULT_AUCTION_SUBTITLE}')"
    )
    parser.add_argument(
        "--tagline",
        default=DEFAULT_LEAGUE_TAG,
        help=f"Tagline pill on cover slide (Default: '{DEFAULT_LEAGUE_TAG}')"
    )
    parser.add_argument(
        "--organizer",
        default=DEFAULT_ORGANIZER,
        help=f"Organizer text on cover slide footer (Default: '{DEFAULT_ORGANIZER}')"
    )
    parser.add_argument(
        "--event-tag",
        default=DEFAULT_EVENT_TAG,
        help=f"Event badge text on cover slide (Default: '{DEFAULT_EVENT_TAG}')"
    )
    args = parser.parse_args()

    data_source = SAMPLE_PLAYERS if args.sample else args.input
    generate_hpl2_presentation(
        players_data=data_source,
        output_path=args.output,
        tagline=args.tagline,
        title=args.title,
        subtitle=args.subtitle,
        event_tag=args.event_tag,
        organizer=args.organizer
    )
