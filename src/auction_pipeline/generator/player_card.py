"""
Player Card Layout Generator
============================
Creates portrait frame, name banners, nationality/rating badges, and metric chips.
"""

import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from src.common.config.settings import (
    THEME_COLOR_CARD_BG, THEME_COLOR_BLUE, THEME_COLOR_GOLD,
    THEME_COLOR_CYAN, THEME_COLOR_WHITE, THEME_COLOR_MUTED
)
from src.common.utils.text_utils import clean_val


def render_player_header_and_photo(slide, row: dict):
    """
    Renders the player portrait photo frame, title header, and quick metric bar.
    """
    # 1. Player Photo Card Frame (Left side)
    photo_frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(1.0), Inches(4.2), Inches(5.5)
    )
    photo_frame.fill.solid()
    photo_frame.fill.fore_color.rgb = THEME_COLOR_CARD_BG
    photo_frame.line.color.rgb = THEME_COLOR_BLUE
    photo_frame.line.width = Pt(1.5)

    img_path = str(row.get("Image_Path", "")).strip()
    img_placed = False
    if img_path and os.path.exists(img_path):
        try:
            slide.shapes.add_picture(
                img_path,
                Inches(0.9), Inches(1.1),
                width=Inches(4.0), height=Inches(5.3)
            )
            img_placed = True
        except Exception:
            img_placed = False

    if not img_placed:
        tf = photo_frame.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "NO PHOTO\nAVAILABLE"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(18)
        p.font.color.rgb = THEME_COLOR_MUTED

    # 2. Header Info Box (Right side)
    player_name = clean_val(row.get("Player_Name"), default="UNKNOWN PLAYER").upper()
    category = clean_val(row.get("Category"), default="GENERAL")
    country = clean_val(row.get("Country"), default="INDIA")
    rating = clean_val(row.get("Rating_Points"), default="-")

    header_box = slide.shapes.add_textbox(Inches(5.3), Inches(0.9), Inches(7.3), Inches(1.1))
    tf_header = header_box.text_frame
    tf_header.word_wrap = True

    p_name = tf_header.paragraphs[0]
    p_name.text = player_name
    p_name.font.size = Pt(32)
    p_name.font.bold = True
    p_name.font.color.rgb = THEME_COLOR_WHITE

    p_sub = tf_header.add_paragraph()
    p_sub.text = f"{category.upper()}  |  {country.upper()}  |  RATING: {rating} / 10"
    p_sub.font.size = Pt(15)
    p_sub.font.bold = True
    p_sub.font.color.rgb = THEME_COLOR_GOLD
    p_sub.space_before = Pt(4)

    # 3. Quick Metric Chips (Matches & Base Price)
    matches_val = clean_val(row.get("Matches"), default="-")
    base_price = clean_val(row.get("Base_Price"), default="-")

    metrics_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.3), Inches(2.2), Inches(7.2), Inches(0.7)
    )
    metrics_shape.fill.solid()
    metrics_shape.fill.fore_color.rgb = THEME_COLOR_CARD_BG
    metrics_shape.line.color.rgb = THEME_COLOR_CYAN
    metrics_shape.line.width = Pt(1)

    tf_m = metrics_shape.text_frame
    tf_m.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_m = tf_m.paragraphs[0]
    p_m.text = f"MATCHES PLAYED: {matches_val}     •     BASE PRICE: {base_price}"
    p_m.alignment = PP_ALIGN.CENTER
    p_m.font.size = Pt(15)
    p_m.font.bold = True
    p_m.font.color.rgb = THEME_COLOR_WHITE
