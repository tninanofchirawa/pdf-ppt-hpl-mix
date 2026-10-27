"""
Presentation Styler & Theme Palette
===================================
Visual constants and shape-styling helper routines for presentation generation.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from src.common.config.settings import (
    SLIDE_WIDTH_16_9, SLIDE_HEIGHT_16_9,
    THEME_COLOR_BG, THEME_COLOR_CARD_BG,
    THEME_COLOR_GOLD, THEME_COLOR_BLUE,
    THEME_COLOR_CYAN, THEME_COLOR_EMERALD,
    THEME_COLOR_WHITE, THEME_COLOR_MUTED,
    THEME_COLOR_TABLE_HEADER, THEME_COLOR_TABLE_ROW_ALT
)


def create_deck_canvas() -> Presentation:
    """
    Initializes a PowerPoint presentation deck configured in 16:9 widescreen format.
    """
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_16_9)
    prs.slide_height = Inches(SLIDE_HEIGHT_16_9)
    return prs


def apply_solid_background(slide, prs: Presentation, color: RGBColor = THEME_COLOR_BG):
    """
    Draws a full-bleed solid background canvas shape.
    """
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg
