"""
Statistics Table Generator
==========================
Creates formatted batting and bowling tables on player auction slides.
"""

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

from src.common.config.settings import (
    THEME_COLOR_CARD_BG, THEME_COLOR_TABLE_HEADER,
    THEME_COLOR_TABLE_ROW_ALT, THEME_COLOR_WHITE, THEME_COLOR_MUTED
)
from src.common.utils.text_utils import clean_val


def render_stats_table(slide, row: dict):
    """
    Renders structured Batting and Bowling statistics table with header formatting.
    """
    table_shape = slide.shapes.add_table(3, 4, Inches(5.3), Inches(3.2), Inches(7.2), Inches(2.6))
    tbl = table_shape.table

    headers = ["BATTING STAT", "VALUE", "BOWLING STAT", "VALUE"]
    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = THEME_COLOR_TABLE_HEADER
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = THEME_COLOR_WHITE

    data_rows = [
        [
            "Batting Avg", clean_val(row.get("Bat_Avg")),
            "Bowling Econ", clean_val(row.get("Bowl_Econ"))
        ],
        [
            "Runs Scored", clean_val(row.get("Runs")),
            "Wickets Taken", clean_val(row.get("Wickets"))
        ]
    ]

    for r_idx, row_vals in enumerate(data_rows):
        for c_idx, val in enumerate(row_vals):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = THEME_COLOR_TABLE_ROW_ALT if r_idx % 2 == 1 else THEME_COLOR_CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(14)
            p.font.bold = (c_idx in [1, 3])
            p.font.color.rgb = THEME_COLOR_WHITE if c_idx in [1, 3] else THEME_COLOR_MUTED
