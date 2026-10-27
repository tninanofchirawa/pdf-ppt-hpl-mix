"""
Divider Slide Generator
=======================
Builds high-impact category title slides and countdown checkpoint slides.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from src.common.config.settings import (
    THEME_COLOR_CARD_BG, THEME_COLOR_GOLD, THEME_COLOR_WHITE
)
from src.auction_pipeline.generator.styler import apply_solid_background


def add_divider_slide(prs: Presentation, title: str, subtitle: str = ""):
    """
    Renders an eye-catching section divider or countdown alert slide.
    """
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_solid_background(slide, prs)

    # Decorative container card
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.5), Inches(1.5),
        prs.slide_width - Inches(3.0), prs.slide_height - Inches(3.0)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = THEME_COLOR_CARD_BG
    card.line.color.rgb = THEME_COLOR_GOLD
    card.line.width = Pt(2)

    # Text box
    tb = slide.shapes.add_textbox(Inches(1.8), Inches(2.2), prs.slide_width - Inches(3.6), Inches(2.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title.upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(50)
    p.font.bold = True
    p.font.color.rgb = THEME_COLOR_WHITE

    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(24)
        p2.font.color.rgb = THEME_COLOR_GOLD
        p2.space_before = Pt(14)
